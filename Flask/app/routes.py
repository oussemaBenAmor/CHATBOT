from flask import Blueprint, request, jsonify, render_template, current_app
import pdfplumber
import openpyxl
from docx import Document
import re
import os
import json
from werkzeug.utils import secure_filename
from dotenv import load_dotenv
import spacy
from typing import List, Dict, Tuple
import glob
from pymongo import MongoClient
import difflib
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
import numpy as np
import dropbox
from sentence_transformers import SentenceTransformer
import torch
import tempfile
from .utils import (
    clean_text, extract_key_phrases, calculate_semantic_similarity,
    find_most_similar_sentences, categorize_transaction_question,
    extract_question_intent, format_answer_for_intent
)
from .web_scraper import create_web_scraper

# Load environment variables
load_dotenv()

# Initialize blueprint
main = Blueprint('main', __name__)

# Allowed file extensions
ALLOWED_EXTENSIONS = {'pdf', 'doc', 'docx', 'xls', 'xlsx', 'txt'}
sbert_model = SentenceTransformer('all-MiniLM-L6-v2')
# Load spaCy model
try:
    nlp = spacy.load("en_core_web_md")
except OSError:
    nlp = spacy.load("en_core_web_sm")
    print("Warning: Using small spaCy model. Install 'en_core_web_md' for better similarity matching.")

nlp.add_pipe("sentencizer")

APP_KEY = os.getenv("DROPBOX_APP_KEY")
APP_SECRET = os.getenv("DROPBOX_APP_SECRET")
REFRESH_TOKEN = os.getenv("DROPBOX_REFRESH_TOKEN")

dbx = dropbox.Dropbox(
    app_key=APP_KEY,
    app_secret=APP_SECRET,
    oauth2_refresh_token=REFRESH_TOKEN
)

# MongoDB Atlas configuration
mongo_client = MongoClient(os.getenv('MONGODB_URI'))
database_name = os.getenv('DATABASE_NAME')
collection_name = os.getenv('COLLECTION_NAME')
database = mongo_client[database_name]
collection = database[collection_name]

# Initialize web scraper
web_scraper = create_web_scraper()

# Transaction type keywords
TRANSACTION_KEYWORDS = {
    'refunds': [
        'refund', 'refunds', 'return', 'returns', 'credit', 'credits', 
        'money back', 'reimbursement', 'repayment', 'cash back', 'chargeback'
    ],
    'payments': [
        'payment', 'payments', 'pay', 'paid', 'card', 'cards', 'bank', 
        'banking', 'cash', 'transfer', 'credit card', 'debit card', 'online payment'
    ],
    'transfers': [
        'transfer', 'transfers', 'move', 'moves', 'send', 'sends', 'wire', 
        'wires', 'bank transfer', 'money transfer', 'electronic transfer'
    ],
    'exchanges': [
        'exchange', 'exchanges', 'swap', 'swaps', 'trade', 'trades', 
        'replace', 'replaces', 'substitution', 'conversion'
    ]
}

def allowed_file(filename: str) -> bool:
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_text_from_pdf(file_path: str) -> str:
    text = ""
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                page_text = re.sub(r'\s+', ' ', page_text).strip()
                text += page_text + "\n\n"
    return text


# def filter_conditions_by_relevance(conditions: list, question_focus: list, question_doc, threshold: float = 0.25) -> list:
#     """Filter a list of conditions for relevance to the question."""
#     return [
#         cond for cond in conditions
#         if calculate_relevance_score(cond, question_focus, question_doc) > threshold
#     ]
def filter_conditions_by_relevance(conditions: list, question_focus: list, question_doc, sbert_model, threshold: float = 0.25) -> list:
    """Filter a list of conditions for relevance to the question using SBERT."""
    return [
        cond for cond in conditions
        if calculate_relevance_score(cond, question_focus, question_doc, sbert_model) > threshold
    ]
def download_dropbox_folder(folder_path="/chatbot-training"):
    """Download all files from Dropbox folder to a temp directory and return local path."""
    tmp_dir = tempfile.mkdtemp()
    result = dbx.files_list_folder(folder_path)
    
    for entry in result.entries:
        if isinstance(entry, dropbox.files.FileMetadata):
            local_path = os.path.join(tmp_dir, entry.name)
            with open(local_path, "wb") as f:
                metadata, res = dbx.files_download(entry.path_lower)
                f.write(res.content)
    
    return tmp_dir
def extract_text_from_excel(file_path: str) -> str:
    text = ""
    workbook = openpyxl.load_workbook(file_path, read_only=True)
    for sheet in workbook:
        for row in sheet.rows:
            row_text = " ".join(str(cell.value) for cell in row if cell.value)
            if row_text:
                text += row_text + "\n"
    return re.sub(r'\s+', ' ', text).strip()

def extract_text_from_word(file_path: str) -> str:
    doc = Document(file_path)
    text = ""
    for para in doc.paragraphs:
        if para.text:
            text += para.text + "\n"
    return re.sub(r'\s+', ' ', text).strip()

def extract_text(file_path: str) -> str:
    extension = file_path.rsplit('.', 1)[1].lower()
    try:
        if extension == 'pdf':
            return extract_text_from_pdf(file_path)
        elif extension in ['xlsx', 'xls']:
            return extract_text_from_excel(file_path)
        elif extension == 'docx':
            return extract_text_from_word(file_path)
        elif extension == 'txt':
            with open(file_path, 'r', encoding='utf-8') as f:
                return re.sub(r'\s+', ' ', f.read()).strip()
        else:
            return ""
    except Exception as e:
        print(f"Error extracting text from {file_path}: {e}")
        return ""

def preprocess_text(text: str) -> List[str]:
    doc = nlp(text)
    sentences = []
    for sent in doc.sents:
        s = sent.text.strip()
        # Skip lines that look like table headers or section titles
        if not s:
            continue
        if re.match(r'^(Summary|Condition Type|Details|Explanation|Requirements?|Procedures?|Restrictions?|Timeframes?|Fees|Section|Contact)\b', s, re.IGNORECASE):
            continue
        if len(s.split()) < 4:  # Skip very short lines
            continue
        sentences.append(s)
    return sentences

def detect_transaction_type(question: str) -> Tuple[str, float]:
    question_lower = question.lower()
    question_doc = nlp(question_lower)
    keyword_scores = categorize_transaction_question(question)
    semantic_scores = {}
    for t_type in TRANSACTION_KEYWORDS.keys():
        type_doc = nlp(t_type)
        semantic_scores[t_type] = question_doc.similarity(type_doc)
    key_phrases = extract_key_phrases(question, nlp)
    phrase_scores = {}
    for t_type, keywords in TRANSACTION_KEYWORDS.items():
        matches = sum(1 for phrase in key_phrases if any(keyword in phrase for keyword in keywords))
        phrase_scores[t_type] = matches / len(key_phrases) if key_phrases else 0
    final_scores = {}
    for t_type in TRANSACTION_KEYWORDS.keys():
        keyword_score = keyword_scores.get(t_type, 0)
        semantic_score = semantic_scores.get(t_type, 0)
        phrase_score = phrase_scores.get(t_type, 0)
        final_score = (keyword_score * 0.5) + (semantic_score * 0.3) + (phrase_score * 0.2)
        final_scores[t_type] = final_score
    if final_scores:
        best_type = max(final_scores, key=final_scores.get)
        best_score = final_scores[best_type]
        if best_score > 0.2:
            return best_type, best_score
    return None, 0.0

def extract_question_focus(question: str) -> List[str]:
    question_doc = nlp(question.lower())
    focus_words = []
    for token in question_doc:
        if token.pos_ in ['NOUN', 'VERB', 'ADJ'] and not token.is_stop and not token.is_punct and len(token.text) > 2:
            focus_words.append(token.lemma_.lower())
    question_words = ['what', 'how', 'when', 'where', 'why', 'which', 'who']
    for word in question_words:
        if word in question.lower():
            focus_words.append(word)
    key_phrases = extract_key_phrases(question, nlp)
    focus_words.extend(key_phrases)
    return list(set(focus_words))

# def calculate_relevance_score(sentence: str, question_focus: List[str], question_doc) -> float:
#     sentence_doc = nlp(sentence.lower())
#     try:
#         vectorizer = TfidfVectorizer(stop_words='english', ngram_range=(1, 2))
#         tfidf_matrix = vectorizer.fit_transform([sentence.lower(), ' '.join(question_focus)])
#         tfidf_similarity = cosine_similarity(tfidf_matrix[0:1], tfidf_matrix[1:2])[0][0]
#     except:
#         tfidf_similarity = 0.0
#     semantic_similarity = question_doc.similarity(sentence_doc)
#     sentence_words = set(sentence.lower().split())
#     keyword_overlap = len(set(question_focus) & sentence_words) / len(question_focus) if question_focus else 0
#     sentence_key_phrases = extract_key_phrases(sentence, nlp)
#     phrase_overlap = len(set(question_focus) & set(sentence_key_phrases)) / len(question_focus) if question_focus else 0
#     # Increase semantic similarity weight
#     return (tfidf_similarity * 0.2 + semantic_similarity * 0.5 + keyword_overlap * 0.15 + phrase_overlap * 0.15)
def calculate_relevance_score(sentence: str, question_focus: List[str], question_doc, sbert_model) -> float:
    # Use SBERT for semantic similarity
    semantic_similarity = calculate_semantic_similarity(' '.join(question_focus), sentence, sbert_model)
    # You can keep TF-IDF and other features if you want, or just use SBERT
    return semantic_similarity
def filter_relevant_sentences(sentences: List[str], question_focus: List[str], question_doc, sbert_model, threshold: float = 0.25) -> List[str]:
    relevant_sentences = []
    for sentence in sentences:
        relevance = calculate_relevance_score(sentence, question_focus, question_doc, sbert_model)
        if relevance > threshold:
            relevant_sentences.append((sentence, relevance))
    relevant_sentences.sort(key=lambda x: x[1], reverse=True)
    return [sent for sent, _ in relevant_sentences[:8]]

def process_urls_in_question(question: str, transaction_type: str) -> List[Dict]:
    try:
        return web_scraper.process_urls_in_question(question, transaction_type)
    except Exception as e:
        print(f"Error processing URLs: {e}")
        return []

def train_chatbot(folder_path="/chatbot-training") -> None:
    training_folder = download_dropbox_folder(folder_path)
    if not os.path.exists(training_folder):
        print(f"Training folder {training_folder} not found.")
        return
    conditions_by_type = {}
    for file_path in glob.glob(os.path.join(training_folder, '*')):
        if allowed_file(file_path):
            base_name = os.path.basename(file_path).lower()
            transaction_type = None
            for t_type in ['refunds', 'payments', 'exchanges', 'transfers']:
                if t_type in base_name:
                    transaction_type = t_type
                    break
            if not transaction_type:
                transaction_type = base_name.rsplit('_v', 1)[0]
            text = extract_text(file_path)
            if text:
                sentences = preprocess_text(text)
                if transaction_type not in conditions_by_type:
                    conditions_by_type[transaction_type] = set()
                conditions_by_type[transaction_type].update(sentences)
                print(f"Processed conditions for {transaction_type} from {file_path}")
            else:
                print(f"Failed to extract text from {file_path}")
    for transaction_type, sentences in conditions_by_type.items():
        unique_id = f"{transaction_type}_consolidated"
        collection.update_one(
            {'id': unique_id},
            {'$set': {'transaction_type': transaction_type, 'conditions': list(sentences)}},
            upsert=True
        )
        print(f"Trained and saved deduplicated conditions for {transaction_type}")

def get_conditions_from_db(transaction_type: str) -> List[str]:
    try:
        items = collection.find({'transaction_type': transaction_type})
        conditions = []
        for item in items:
            conditions.extend(item.get('conditions', []))
        return list(dict.fromkeys(conditions))
    except Exception as e:
        print(f"Error querying MongoDB for {transaction_type}: {e}")
        return []

# def generate_focused_answer(question: str, relevant_sentences: List[str], transaction_type: str, url_data: List[Dict] = None, file_content: str = None) -> str:
#     if file_content:
#         answer = f"📄 **Information from uploaded file about {transaction_type}:**\n\n"
#         sentences = re.split(r'[.!?]+', file_content)
#         relevant_sentences = [s.strip() for s in sentences if len(s.strip()) > 20]
#         transaction_conditions = web_scraper.extract_transaction_conditions(file_content, transaction_type)
#         if transaction_conditions.get('requirements'):
#             answer += "**Requirements:**\n"
#             for i, req in enumerate(transaction_conditions['requirements'][:5], 1):
#                 answer += f"{i}. {req}\n"
#             answer += "\n"
#         if transaction_conditions.get('procedures'):
#             answer += "**Procedures:**\n"
#             for i, proc in enumerate(transaction_conditions['procedures'][:5], 1):
#                 answer += f"{i}. {proc}\n"
#             answer += "\n"
#         if transaction_conditions.get('restrictions'):
#             answer += "**Restrictions:**\n"
#             for i, restriction in enumerate(transaction_conditions['restrictions'][:5], 1):
#                 answer += f"{i}. {restriction}\n"
#             answer += "\n"
#         if transaction_conditions.get('timeframes'):
#             answer += "**Timeframes:**\n"
#             for i, time in enumerate(transaction_conditions['timeframes'][:5], 1):
#                 answer += f"{i}. {time}\n"
#             answer += "\n"
#         if transaction_conditions.get('fees'):
#             answer += "**Fees & Costs:**\n"
#             for i, fee in enumerate(transaction_conditions['fees'][:5], 1):
#                 answer += f"{i}. {fee}\n"
#             answer += "\n"
#         if relevant_sentences and not any(transaction_conditions.values()):
#             answer += "**Key Information:**\n"
#             for i, sentence in enumerate(relevant_sentences[:10], 1):
#                 answer += f"{i}. {sentence}\n"
#         return answer
#     if url_data and any(d['status'] == 'success' for d in url_data):
#         answer = f"📎 **Information from websites about {transaction_type}:**\n\n"
#         for url_info in url_data:
#             if url_info['status'] == 'success':
#                 answer += f"**From: {url_info['title']}**\n"
#                 answer += f"**URL:** {url_info['url']}\n\n"
#                 if url_info.get('raw_content'):
#                     content = url_info['raw_content']
#                     sentences = re.split(r'[.!?]+', content)
#                     relevant_sentences_from_url = [s.strip() for s in sentences if len(s.strip()) > 20][:15]
#                     answer += "**Key Information:**\n"
#                     for i, sentence in enumerate(relevant_sentences_from_url, 1):
#                         answer += f"{i}. {sentence}\n"
#                     answer += "\n" + "─" * 50 + "\n\n"
#                 elif 'transaction_conditions' in url_info and url_info['transaction_conditions']:
#                     conditions = url_info['transaction_conditions']
#                     if conditions.get('requirements'):
#                         answer += "**Requirements:**\n"
#                         for i, req in enumerate(conditions['requirements'][:5], 1):
#                             answer += f"{i}. {req}\n"
#                         answer += "\n"
#                     if conditions.get('procedures'):
#                         answer += "**Procedures:**\n"
#                         for i, proc in enumerate(conditions['procedures'][:5], 1):
#                             answer += f"{i}. {proc}\n"
#                         answer += "\n"
#                     if conditions.get('restrictions'):
#                         answer += "**Restrictions:**\n"
#                         for i, restriction in enumerate(conditions['restrictions'][:5], 1):
#                             answer += f"{i}. {restriction}\n"
#                         answer += "\n"
#                     if conditions.get('timeframes'):
#                         answer += "**Timeframes:**\n"
#                         for i, time in enumerate(conditions['timeframes'][:5], 1):
#                             answer += f"{i}. {time}\n"
#                         answer += "\n"
#                     if conditions.get('fees'):
#                         answer += "**Fees & Costs:**\n"
#                         for i, fee in enumerate(conditions['fees'][:5], 1):
#                             answer += f"{i}. {fee}\n"
#                         answer += "\n"
#                     answer += "─" * 50 + "\n\n"
#         return answer
#     if not relevant_sentences:
#         return f"No specific information found about {question.lower()} for {transaction_type}."
#     answer = f"Here's what I found about {transaction_type}:\n\n"
#     for i, sentence in enumerate(relevant_sentences[:5], 1):
#         answer += f"{i}. {sentence.strip()}\n"
#     return answer
# def generate_focused_answer(
#     question: str,
#     relevant_sentences: List[str],
#     transaction_type: str,
#     url_data: List[Dict] = None,
#     file_content: str = None,
#     question_focus: List[str] = None,
#     question_doc = None
# ) -> str:
#     if file_content:
#         answer = f"📄 **Information from uploaded file about {transaction_type}:**\n\n"
#         transaction_conditions = web_scraper.extract_transaction_conditions(file_content, transaction_type)
#         # Filter each condition type for relevance if question_focus and question_doc are provided
#         if question_focus is not None and question_doc is not None:
#             filtered_conditions = {
#                 key: filter_conditions_by_relevance(val, question_focus, question_doc)
#                 for key, val in transaction_conditions.items() if isinstance(val, list)
#             }
#         else:
#             filtered_conditions = transaction_conditions

#         for key in ['requirements', 'procedures', 'restrictions', 'timeframes', 'fees']:
#             items = filtered_conditions.get(key, [])
#             if items:
#                 answer += f"**{key.capitalize()}:**\n"
#                 for i, item in enumerate(items[:5], 1):
#                     answer += f"{i}. {item}\n"
#                 answer += "\n"
#         # Fallback to relevant sentences if no relevant conditions found
#         if relevant_sentences and not any(filtered_conditions.values()):
#             answer += "**Key Information:**\n"
#             for i, sentence in enumerate(relevant_sentences[:10], 1):
#                 answer += f"{i}. {sentence}\n"
#         return answer

#     if url_data and any(d['status'] == 'success' for d in url_data):
#         answer = f"📎 **Information from websites about {transaction_type}:**\n\n"
#         for url_info in url_data:
#             if url_info['status'] == 'success':
#                 answer += f"**From: {url_info['title']}**\n"
#                 answer += f"**URL:** {url_info['url']}\n\n"
#                 if url_info.get('raw_content'):
#                     content = url_info['raw_content']
#                     sentences = re.split(r'[.!?]+', content)
#                     relevant_sentences_from_url = [s.strip() for s in sentences if len(s.strip()) > 20][:15]
#                     answer += "**Key Information:**\n"
#                     for i, sentence in enumerate(relevant_sentences_from_url, 1):
#                         answer += f"{i}. {sentence}\n"
#                     answer += "\n" + "─" * 50 + "\n\n"
#                 elif 'transaction_conditions' in url_info and url_info['transaction_conditions']:
#                     conditions = url_info['transaction_conditions']
#                     for key in ['requirements', 'procedures', 'restrictions', 'timeframes', 'fees']:
#                         items = conditions.get(key, [])
#                         if items:
#                             answer += f"**{key.capitalize()}:**\n"
#                             for i, item in enumerate(items[:5], 1):
#                                 answer += f"{i}. {item}\n"
#                             answer += "\n"
#                     answer += "─" * 50 + "\n\n"
#         return answer

#     if not relevant_sentences:
#         return f"No specific information found about {question.lower()} for {transaction_type}."
#     answer = f"Here's what I found about {transaction_type}:\n\n"
#     for i, sentence in enumerate(relevant_sentences[:5], 1):
#         answer += f"{i}. {sentence.strip()}\n"
#     return answer



# def generate_focused_answer(
#     question: str,
#     relevant_sentences: List[str],
#     transaction_type: str,
#     url_data: List[Dict] = None,
#     file_content: str = None,
#     question_focus: List[str] = None,
#     question_doc = None
# ) -> str:
#     def dedup_and_clean(sentences: List[str]) -> List[str]:
#         seen = set()
#         cleaned = []
#         for s in sentences:
#             s = s.strip()
#             # Skip table headers/section titles
#             if re.match(r'^(Summary|Condition Type|Details|Explanation|Requirements?|Procedures?|Restrictions?|Timeframes?|Fees|Section|Contact)\b', s, re.IGNORECASE):
#                  continue
#             if len(s.split()) < 4:
#                 continue
#             if s and s not in seen:
#                 cleaned.append(s)
#                 seen.add(s)
#         return cleaned

#     # Helper to group sentences by keywords
#     def group_sentences(sentences: List[str],sbert_model) -> Dict[str, List[str]]:
#         groups = {"Refundability": [], "Processing": [], "Requirements": [], "Other": []}
#         for s in sentences:
#             s_lower = s.lower()
#             if "refundable" in s_lower or "refund" in s_lower:
#                 groups["Refundability"].append(s)
#             elif "process" in s_lower or "processed" in s_lower or "time" in s_lower or "day" in s_lower:
#                 groups["Processing"].append(s)
#             elif "require" in s_lower or "condition" in s_lower or "eligible" in s_lower:
#                 groups["Requirements"].append(s)
#             else:
#                 groups["Other"].append(s)
#         return groups

#     # If file content, use the existing logic
#     if file_content:
#         answer = f"📄 **Information from uploaded file about {transaction_type}:**\n\n"
#         transaction_conditions = web_scraper.extract_transaction_conditions(file_content, transaction_type)
#         if question_focus is not None and question_doc is not None:
#             filtered_conditions = {
#                 key: filter_conditions_by_relevance(val, question_focus, question_doc,sbert_model)
#                 for key, val in transaction_conditions.items() if isinstance(val, list)
#             }
#         else:
#             filtered_conditions = transaction_conditions

#         for key in ['requirements', 'procedures', 'restrictions', 'timeframes', 'fees']:
#             items = filtered_conditions.get(key, [])
#             if items:
#                 answer += f"**{key.capitalize()}:**\n"
#                 for i, item in enumerate(items[:5], 1):
#                     answer += f"{i}. {item}\n"
#                 answer += "\n"
#         if relevant_sentences and not any(filtered_conditions.values()):
#             answer += "**Key Information:**\n"
#             for i, sentence in enumerate(relevant_sentences[:10], 1):
#                 answer += f"{i}. {sentence}\n"
#         return answer

def generate_focused_answer(
    question: str,
    relevant_sentences: List[str],
    transaction_type: str,
    url_data: List[Dict] = None,
    file_content: str = None,
    question_focus: List[str] = None,
    question_doc = None
) -> str:
    def dedup_and_clean(sentences: List[str]) -> List[str]:
        seen = set()
        cleaned = []
        for s in sentences:
            s = s.strip()
            # Skip table headers/section titles
            if re.match(r'^(Summary|Condition Type|Details|Explanation|Requirements?|Procedures?|Restrictions?|Timeframes?|Fees|Section|Contact)\b', s, re.IGNORECASE):
                 continue
            if len(s.split()) < 4:
                continue
            if s and s not in seen:
                cleaned.append(s)
                seen.add(s)
        return cleaned

    # Helper to group sentences by keywords
    def group_sentences(sentences: List[str], sbert_model) -> Dict[str, List[str]]:
        groups = {"Refundability": [], "Processing": [], "Requirements": [], "Other": []}
        for s in sentences:
            s_lower = s.lower()
            if "refundable" in s_lower or "refund" in s_lower:
                groups["Refundability"].append(s)
            elif "process" in s_lower or "processed" in s_lower or "time" in s_lower or "day" in s_lower:
                groups["Processing"].append(s)
            elif "require" in s_lower or "condition" in s_lower or "eligible" in s_lower:
                groups["Requirements"].append(s)
            else:
                groups["Other"].append(s)
        return groups

    # If file content, use the existing logic
    if file_content:
        answer = f"📄 **Information from uploaded file about {transaction_type}:**\n\n"
        transaction_conditions = web_scraper.extract_transaction_conditions(file_content, transaction_type)
        if question_focus is not None and question_doc is not None:
            filtered_conditions = {
                key: filter_conditions_by_relevance(val, question_focus, question_doc, sbert_model)
                for key, val in transaction_conditions.items() if isinstance(val, list)
            }
        else:
            filtered_conditions = transaction_conditions

        for key in ['requirements', 'procedures', 'restrictions', 'timeframes', 'fees']:
            items = filtered_conditions.get(key, [])
            if items:
                answer += f"**{key.capitalize()}:**\n"
                for i, item in enumerate(items[:5], 1):
                    answer += f"{i}. {item}\n"
                answer += "\n"
        if relevant_sentences and not any(filtered_conditions.values()):
            answer += "**Key Information:**\n"
            for i, sentence in enumerate(relevant_sentences[:10], 1):
                answer += f"{i}. {sentence}\n"
        return answer

    # If web data, use the existing logic
    if url_data and any(d['status'] == 'success' for d in url_data):
        answer = f"📎 **Information from websites about {transaction_type}:**\n\n"
        for url_info in url_data:
            if url_info['status'] == 'success':
                answer += f"**From: {url_info['title']}**\n"
                answer += f"**URL:** {url_info['url']}\n\n"
                if 'transaction_conditions' in url_info and url_info['transaction_conditions']:
                    conditions = url_info['transaction_conditions']
                    for key in ['requirements', 'procedures', 'restrictions', 'timeframes', 'fees', 'general_info']:
                        items = conditions.get(key, [])
                        if items:
                            answer += f"**{key.capitalize()}:**\n"
                            for i, item in enumerate(items, 1):
                                answer += f"{i}. {item}\n"
                            answer += "\n"
                answer += "─" * 50 + "\n\n"
        return answer

    # For DB or fallback, organize and deduplicate
    if not relevant_sentences:
        return f"No specific information found about {question.lower()} for {transaction_type}."

    deduped = dedup_and_clean(relevant_sentences)
    groups = group_sentences(deduped, sbert_model)

    answer = f"**Here's what I found about {transaction_type} (organized):**\n\n"
    for section, items in groups.items():
        if items:
            answer += f"**{section}:**\n"
            for i, item in enumerate(items, 1):
                answer += f"{i}. {item}\n"
            answer += "\n"
    return answer.strip()



    # If web data, use the existing logic
    # if url_data and any(d['status'] == 'success' for d in url_data):
    #     answer = f"📎 **Information from websites about {transaction_type}:**\n\n"
    #     for url_info in url_data:
    #         if url_info['status'] == 'success':
    #             answer += f"**From: {url_info['title']}**\n"
    #             answer += f"**URL:** {url_info['url']}\n\n"
    #             if url_info.get('raw_content'):
    #                 content = url_info['raw_content']
    #                 sentences = re.split(r'[.!?]+', content)
    #                 relevant_sentences_from_url = [s.strip() for s in sentences if len(s.strip()) > 20][:15]
    #                 relevant_sentences_from_url = dedup_and_clean(relevant_sentences_from_url)
    #                 answer += "**Key Information:**\n"
    #                 for i, sentence in enumerate(relevant_sentences_from_url, 1):
    #                     answer += f"{i}. {sentence}\n"
    #                 answer += "\n" + "─" * 50 + "\n\n"
    #             elif 'transaction_conditions' in url_info and url_info['transaction_conditions']:
    #                 conditions = url_info['transaction_conditions']
    #                 for key in ['requirements', 'procedures', 'restrictions', 'timeframes', 'fees']:
    #                     items = conditions.get(key, [])
    #                     if items:
    #                         answer += f"**{key.capitalize()}:**\n"
    #                         for i, item in enumerate(items[:5], 1):
    #                             answer += f"{i}. {item}\n"
    #                         answer += "\n"
    #                 answer += "─" * 50 + "\n\n"
    #     return answer
    if url_data and any(d['status'] == 'success' for d in url_data):
        answer = f"📎 **Information from websites about {transaction_type}:**\n\n"
    for url_info in url_data:
        if url_info['status'] == 'success':
            answer += f"**From: {url_info['title']}**\n"
            answer += f"**URL:** {url_info['url']}\n\n"
            if 'transaction_conditions' in url_info and url_info['transaction_conditions']:
                conditions = url_info['transaction_conditions']
                for key in ['requirements', 'procedures', 'restrictions', 'timeframes', 'fees', 'general_info']:
                    items = conditions.get(key, [])
                    if items:
                        answer += f"**{key.capitalize()}:**\n"
                        for i, item in enumerate(items, 1):
                            answer += f"{i}. {item}\n"
                        answer += "\n"
            answer += "─" * 50 + "\n\n"
    return answer

    # For DB or fallback, organize and deduplicate
    if not relevant_sentences:
        return f"No specific information found about {question.lower()} for {transaction_type}."

    deduped = dedup_and_clean(relevant_sentences)
    groups = group_sentences(deduped, sbert_model)

    answer = f"**Here's what I found about {transaction_type} (organized):**\n\n"
    for section, items in groups.items():
        if items:
            answer += f"**{section}:**\n"
            for i, item in enumerate(items, 1):
                answer += f"{i}. {item}\n"
            answer += "\n"
    return answer.strip()
@main.route('/')
def index():
    return render_template('index.html')

@main.route('/train', methods=['POST'])
def trigger_training():
    training_folder = os.path.join(current_app.config.get('TRAINING_FOLDER', 'training_data'))
    train_chatbot(training_folder)
    return jsonify({'message': 'Training completed and conditions saved to MongoDB Atlas.'})

# @main.route('/chat', methods=['POST'])
# def chat():
#     question = request.form.get('question', '').strip()
#     file = request.files.get('file')
#     if not question and not file:
#         return jsonify({'answer': 'Please provide a question or upload a file.'}), 400

#     question = clean_text(question) if question else ""
#     transaction_type, confidence_score = detect_transaction_type(question) if question else (None, 0.0)
    
#     if not transaction_type and file:
#         transaction_type = 'refunds'  # Default for file uploads if no question
#         confidence_score = 1.0
#     elif not transaction_type:
#         return jsonify({'answer': 'Could not identify a transaction type. Please include terms like "refunds", "payments", "transfers", or "exchanges".'}), 400

#     question_focus = extract_question_focus(question) if question else []
#     question_doc = nlp(question.lower()) if question else nlp("")
    
#     file_content = None
#     file_path = None
#     if file and allowed_file(file.filename):
#         filename = secure_filename(file.filename)
#         file_path = os.path.join(current_app.config.get('UPLOAD_FOLDER', 'uploads'), filename)
#         os.makedirs(os.path.dirname(file_path), exist_ok=True)
#         file.save(file_path)
#         file_content = extract_text(file_path)
#         os.remove(file_path)  # Clean up after processing
#         if not file_content:
#             return jsonify({'answer': f"Failed to extract content from {filename}."}), 400

#     urls_processed = []
#     if question and web_scraper:
#         urls = web_scraper.extract_urls_from_text(question)
#         if urls:
#             urls_processed = web_scraper.process_urls_in_question(question, transaction_type)
#             print(f"Processed {len(urls_processed)} URLs: {urls}")
    
#     if file_content:
#         relevant_sentences = []
#         if question:
#             sentences = preprocess_text(file_content)
#             relevant_sentences = filter_relevant_sentences(sentences, question_focus, question_doc)
#         answer = generate_focused_answer(question, relevant_sentences, transaction_type, file_content=file_content)
#         return jsonify({
#             'answer': answer,
#             'transaction_type': transaction_type,
#             'confidence': confidence_score,
#             'sentences_count': len(relevant_sentences),
#             'source': 'file_upload'
#         })
    
#     if urls_processed and any(url_info['status'] == 'success' for url_info in urls_processed):
#         relevant_sentences = []
#         for url_info in urls_processed:
#             if url_info['status'] == 'success':
#                 content = url_info.get('raw_content', '')
#                 if content:
#                     sentences = re.split(r'[.!?]+', content)
#                     relevant_sentences.extend([s.strip() for s in sentences if len(s.strip()) > 20])
#         relevant_sentences = list(set(relevant_sentences))[:10]
#         answer = generate_focused_answer(question, relevant_sentences, transaction_type, urls_processed)
#         return jsonify({
#             'answer': answer,
#             'transaction_type': transaction_type,
#             'confidence': confidence_score,
#             'sentences_count': len(relevant_sentences),
#             'urls_processed': len(urls_processed),
#             'source': 'web_scraping'
#         })
    
#     all_sentences = get_conditions_from_db(transaction_type)
#     relevant_sentences = find_most_similar_sentences(question, all_sentences, nlp, top_k=5)
#     answer = generate_focused_answer(question, [sentence for sentence, _ in relevant_sentences], transaction_type, urls_processed)
    
#     return jsonify({
#         'answer': answer,
#         'transaction_type': transaction_type,
#         'confidence': confidence_score,
#         'relevant_sentences_count': len(relevant_sentences),
#         'question_focus': question_focus[:5],
#         'urls_processed': len(urls_processed),
#         'source': 'database'
#     })
@main.route('/chat', methods=['POST'])
def chat():
    question = request.form.get('question', '').strip()
    file = request.files.get('file')
    if not question and not file:
        return jsonify({'answer': 'Please provide a question or upload a file.'}), 400

    question = clean_text(question) if question else ""
    transaction_type, confidence_score = detect_transaction_type(question) if question else (None, 0.0)
    
    if not transaction_type and file:
        transaction_type = 'refunds'  # Default for file uploads if no question
        confidence_score = 1.0
    elif not transaction_type:
        return jsonify({'answer': 'Could not identify a transaction type. Please include terms like "refunds", "payments", "transfers", or "exchanges".'}), 400

    question_focus = extract_question_focus(question) if question else []
    question_doc = nlp(question.lower()) if question else nlp("")
    
    file_content = None
    file_path = None
    if file and allowed_file(file.filename):
        filename = secure_filename(file.filename)
        file_path = os.path.join(current_app.config.get('UPLOAD_FOLDER', 'uploads'), filename)
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        file.save(file_path)
        file_content = extract_text(file_path)
        os.remove(file_path)  # Clean up after processing
        if not file_content:
            return jsonify({'answer': f"Failed to extract content from {filename}."}), 400

    urls_processed = []
    if question and web_scraper:
        urls = web_scraper.extract_urls_from_text(question)
        if urls:
            urls_processed = web_scraper.process_urls_in_question(question, transaction_type)
            print(f"Processed {len(urls_processed)} URLs: {urls}")
    
    if file_content:
        relevant_sentences = []
        if question:
            sentences = preprocess_text(file_content)
            relevant_sentences = filter_relevant_sentences(sentences, question_focus, question_doc,sbert_model)
        answer = generate_focused_answer(
            question,
            relevant_sentences,
            transaction_type,
            file_content=file_content,
            question_focus=question_focus,
            question_doc=question_doc
        )
        return jsonify({
            'answer': answer,
            'transaction_type': transaction_type,
            'confidence': confidence_score,
            'sentences_count': len(relevant_sentences),
            'source': 'file_upload'
        })
    
    if urls_processed and any(url_info['status'] == 'success' for url_info in urls_processed):
        relevant_sentences = []
        for url_info in urls_processed:
            if url_info['status'] == 'success':
                content = url_info.get('raw_content', '')
                if content:
                    sentences = re.split(r'[.!?]+', content)
                    relevant_sentences.extend([s.strip() for s in sentences if len(s.strip()) > 20])
        relevant_sentences = list(set(relevant_sentences))[:10]
        answer = generate_focused_answer(question, relevant_sentences, transaction_type, urls_processed)
        return jsonify({
            'answer': answer,
            'transaction_type': transaction_type,
            'confidence': confidence_score,
            'sentences_count': len(relevant_sentences),
            'urls_processed': len(urls_processed),
            'source': 'web_scraping'
        })
    
    all_sentences = get_conditions_from_db(transaction_type)
    relevant_sentences = find_most_similar_sentences(question, all_sentences, sbert_model, top_k=5)
    answer = generate_focused_answer(question, [sentence for sentence, _ in relevant_sentences], transaction_type, urls_processed)
    
    return jsonify({
        'answer': answer,
        'transaction_type': transaction_type,
        'confidence': confidence_score,
        'relevant_sentences_count': len(relevant_sentences),
        'question_focus': question_focus[:5],
        'urls_processed': len(urls_processed),
        'source': 'database'
    })

@main.route('/reload_training', methods=['POST'])
def reload_training():
    return jsonify({'message': 'Training is now manually triggered via /train endpoint.'})